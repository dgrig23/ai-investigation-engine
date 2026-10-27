"""File loaders: turn raw bytes of a supported format into normalized text + metadata.

Every loader produces "light markdown": headings become `#` lines and paragraphs are
separated by blank lines. That shared shape lets one structure-aware chunker work for
every format. Supported formats: PDF, DOCX, Markdown, HTML, TXT, CSV, JSON, JSONL.
"""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


class UnsupportedFormatError(ValueError):
    pass


class DocumentParseError(ValueError):
    pass


@dataclass
class LoadedFile:
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)
    title: str = ""


def normalize_text(text: str) -> str:
    """Unicode-normalize and tidy whitespace without destroying paragraph breaks."""
    text = unicodedata.normalize("NFC", text).replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    text = re.sub(r"[  ]+", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ---------------------------------------------------------------- helpers

_FRONTMATTER_RE = re.compile(r"\A---\n(.*?)\n---\n", re.S)


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    """Parse a simple `key: value` YAML-style front matter block at the top of a file."""
    match = _FRONTMATTER_RE.match(text)
    if not match:
        return {}, text
    meta: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip().lower()] = value.strip().strip('"').strip("'")
    return meta, text[match.end():]


def _first_heading(text: str) -> str:
    for line in text.splitlines():
        if line.startswith("#"):
            return line.lstrip("#").strip()
    return ""


def _looks_like_heading(line: str) -> bool:
    """Heuristic for PDFs, which have no structural markup: short, no final period, title-like."""
    if not line or len(line) > 70 or line.endswith((".", ",", ";", ":")):
        return False
    if re.match(r"^\d+(\.\d+)*\.?\s+[A-Z]", line):  # "2. Root cause", "3.1 Timeline"
        return True
    words = line.split()
    if len(words) > 8:
        return False
    capitalized = sum(1 for w in words if w[:1].isupper() or not w[:1].isalpha())
    return capitalized == len(words) and len(words) <= 6


# ---------------------------------------------------------------- loaders


def load_pdf(data: bytes) -> LoadedFile:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises many exception types for broken files
        raise DocumentParseError(f"Could not read PDF: {exc}") from exc
    lines_out: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        for line in page_text.splitlines():
            line = line.strip()
            if not line:
                lines_out.append("")
            elif _looks_like_heading(line):
                lines_out += ["", f"## {line}", ""]
            else:
                lines_out.append(line)
        lines_out.append("")
    meta: dict[str, Any] = {"pages": len(reader.pages)}
    info = reader.metadata or {}
    for src, dst in (("/Title", "title"), ("/Author", "author"), ("/Subject", "doc_type"), ("/Keywords", "keywords")):
        if info.get(src):
            meta[dst] = str(info[src])
    if info.get("/CreationDate"):
        raw = str(info["/CreationDate"])  # e.g. D:20260318120000
        m = re.search(r"(\d{4})(\d{2})(\d{2})", raw)
        if m:
            meta["date"] = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    text = _join_pdf_lines(lines_out)
    return LoadedFile(text=text, metadata=meta, title=meta.get("title", ""))


def _join_pdf_lines(lines: list[str]) -> str:
    """PDF text breaks lines at the page width; re-join wrapped lines into paragraphs."""
    paragraphs: list[str] = []
    current: list[str] = []
    for line in lines:
        if not line or line.startswith("#") or line.startswith(("- ", "• ")):
            if current:
                paragraphs.append(" ".join(current))
                current = []
            if line:
                paragraphs.append(line)
        else:
            current.append(line)
    if current:
        paragraphs.append(" ".join(current))
    return "\n\n".join(paragraphs)


def load_docx(data: bytes) -> LoadedFile:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(f"Could not read DOCX: {exc}") from exc
    parts: list[str] = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading"):
            level = int(re.sub(r"\D", "", style) or 1)
            parts.append("#" * min(level, 6) + " " + text)
        elif style == "title":
            parts.append("# " + text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    props = document.core_properties
    meta: dict[str, Any] = {}
    if props.author:
        meta["author"] = props.author
    if props.category:
        meta["doc_type"] = props.category
    if props.keywords:
        meta["keywords"] = props.keywords
    if props.created:
        meta["date"] = props.created.date().isoformat()
    return LoadedFile(text="\n\n".join(parts), metadata=meta, title=props.title or "")


def load_markdown(data: bytes) -> LoadedFile:
    raw = data.decode("utf-8", errors="replace")
    meta, body = parse_frontmatter(raw.replace("\r\n", "\n"))
    return LoadedFile(text=body, metadata=dict(meta), title=meta.get("title", "") or _first_heading(body))


def load_txt(data: bytes) -> LoadedFile:
    return LoadedFile(text=data.decode("utf-8", errors="replace"))


def load_html(data: bytes) -> LoadedFile:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(data, "html.parser")
    meta: dict[str, Any] = {}
    for tag in soup.find_all("meta"):
        name = (tag.get("name") or tag.get("property") or "").lower()
        if name in {"author", "date", "doc_type", "keywords", "description"} and tag.get("content"):
            meta[name] = tag["content"]
    title = soup.title.get_text(strip=True) if soup.title else ""
    for bad in soup(["script", "style", "nav", "footer", "noscript"]):
        bad.decompose()
    parts: list[str] = []
    body = soup.body or soup
    for el in body.find_all(["h1", "h2", "h3", "h4", "p", "li", "tr", "pre", "blockquote"]):
        if el.name == "tr":
            cells = [c.get_text(" ", strip=True) for c in el.find_all(["td", "th"])]
            text = " | ".join(cells)
        else:
            text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in {"h1", "h2", "h3", "h4"}:
            parts.append("#" * int(el.name[1]) + " " + text)
        elif el.name == "li":
            parts.append("- " + text)
        else:
            parts.append(text)
    return LoadedFile(text="\n\n".join(parts), metadata=meta, title=title)


def _record_to_text(record: dict[str, Any]) -> str:
    """Flatten one record (CSV row / JSON object) to `key: value; key: value`."""
    if {"user", "text"} <= record.keys():  # chat-style export (e.g. Slack)
        ts = record.get("ts") or record.get("timestamp") or ""
        return f"[{ts}] {record['user']}: {record['text']}".strip()
    items = []
    for key, value in record.items():
        if isinstance(value, dict | list):
            value = json.dumps(value, ensure_ascii=False)
        items.append(f"{key}: {value}")
    return "; ".join(items)


def load_csv(data: bytes) -> LoadedFile:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        rows = list(csv.DictReader(io.StringIO(text)))
    except csv.Error as exc:
        raise DocumentParseError(f"Could not read CSV: {exc}") from exc
    lines = [_record_to_text(row) for row in rows if any((v or "").strip() for v in row.values())]
    return LoadedFile(text="\n\n".join(lines), metadata={"records": len(lines), "record_based": True})


def _find_records(obj: Any) -> list[dict[str, Any]] | None:
    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        return obj
    if isinstance(obj, dict):
        for value in obj.values():
            found = _find_records(value)
            if found:
                return found
    return None


def load_json(data: bytes) -> LoadedFile:
    try:
        obj = json.loads(data.decode("utf-8", errors="replace"))
    except json.JSONDecodeError as exc:
        raise DocumentParseError(f"Invalid JSON: {exc}") from exc
    meta: dict[str, Any] = {"record_based": True}
    title = ""
    if isinstance(obj, dict):
        title = str(obj.get("title") or obj.get("channel") or obj.get("name") or "")
        for key in ("doc_type", "date", "author"):
            if isinstance(obj.get(key), str):
                meta[key] = obj[key]
    records = _find_records(obj)
    if records is None:
        return LoadedFile(text=json.dumps(obj, indent=1, ensure_ascii=False), metadata=meta, title=title)
    meta["records"] = len(records)
    return LoadedFile(text="\n\n".join(_record_to_text(r) for r in records), metadata=meta, title=title)


def load_jsonl(data: bytes) -> LoadedFile:
    records = []
    for i, line in enumerate(data.decode("utf-8", errors="replace").splitlines()):
        if line.strip():
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as exc:
                raise DocumentParseError(f"Invalid JSON on line {i + 1}: {exc}") from exc
    lines = [_record_to_text(r) if isinstance(r, dict) else str(r) for r in records]
    return LoadedFile(text="\n\n".join(lines), metadata={"records": len(lines), "record_based": True})


LOADERS: dict[str, Callable[[bytes], LoadedFile]] = {
    "pdf": load_pdf,
    "docx": load_docx,
    "md": load_markdown,
    "markdown": load_markdown,
    "txt": load_txt,
    "html": load_html,
    "htm": load_html,
    "csv": load_csv,
    "json": load_json,
    "jsonl": load_jsonl,
}

SUPPORTED_EXTENSIONS = sorted(LOADERS)


def detect_format(filename: str) -> str:
    ext = Path(filename).suffix.lower().lstrip(".")
    if ext not in LOADERS:
        raise UnsupportedFormatError(f"Unsupported file type '.{ext}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}")
    return {"markdown": "md", "htm": "html"}.get(ext, ext)


def load_file(filename: str, data: bytes) -> LoadedFile:
    """Parse any supported file. Raises UnsupportedFormatError / DocumentParseError."""
    fmt = detect_format(filename)
    loaded = LOADERS[fmt](data)
    loaded.text = normalize_text(loaded.text)
    if not loaded.text:
        raise DocumentParseError(f"No text could be extracted from '{filename}'")
    loaded.metadata["format"] = fmt
    if not loaded.title:
        loaded.title = _first_heading(loaded.text) or Path(filename).stem.replace("_", " ").replace("-", " ")
    return loaded
